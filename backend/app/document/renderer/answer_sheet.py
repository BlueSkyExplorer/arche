"""Deterministic renderer for reviewed answer-sheet semantic snapshots."""

from __future__ import annotations

import math
import re
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from io import BytesIO
from uuid import UUID

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import _Cell

from app.document.ooxml import append_page_number, normalize_zip_timestamps, set_table_borders
from app.document.renderer import RenderTemplateProfile, _configure_styles, _format_marks
from app.exam.extraction.answer_sheet import (
    AnsNode,
    AnsSheet,
    AnswerContent,
    AnswerImage,
    AnswerParagraph,
    AnswerTable,
    UnsupportedAnswerContent,
)
from app.schemas.answer_sheet_export import DocumentMetadata

AnswerAssetResolver = Callable[[str], bytes]
LogoAssetResolver = Callable[[UUID], bytes]


@dataclass(frozen=True)
class AnswerSheetRenderResult:
    docx: bytes | None
    validation: dict[str, object]


def _substitute(value: str, metadata: DocumentMetadata) -> str:
    values = metadata.model_dump()
    return re.sub(
        r"{{\s*([a-z_]+)\s*}}",
        lambda match: str(values.get(match.group(1), match.group(0))),
        value,
    ).strip()


def _effective_content(node: AnsNode) -> list[AnswerContent]:
    if node.answer_content:
        return node.answer_content
    return [AnswerParagraph(text=line) for line in node.answer]


def validate_answer_sheet_render(
    sheet: AnsSheet, available_assets: set[str]
) -> dict[str, object]:
    issues: list[dict[str, str]] = []
    stats = {
        "sections": len(sheet.sections),
        "mcq_items": 0,
        "question_nodes": 0,
        "leaves": 0,
        "paragraphs": 0,
        "tables": 0,
        "images": 0,
    }
    for warning in sheet.warnings:
        issues.append(
            {
                "code": warning.code,
                "severity": "warning",
                "message": warning.message,
                "path": warning.path or warning.section or "answer_sheet",
            }
        )
    for section_index, section in enumerate(sheet.sections):
        stats["mcq_items"] += len(section.mcq or [])
        if section.standalone_non_text:
            issues.append(
                {
                    "code": "standalone_non_text",
                    "severity": "blocking",
                    "message": "Detached non-text content cannot be assigned to an answer node.",
                    "path": f"sections[{section_index}]",
                }
            )

        def walk(node: AnsNode, path: str) -> None:
            stats["question_nodes"] += 1
            if node.children:
                if node.marks is not None:
                    issues.append(
                        {
                            "code": "parent_marks_not_authoritative",
                            "severity": "blocking",
                            "message": "A parent node cannot carry authoritative marks.",
                            "path": path,
                        }
                    )
            else:
                stats["leaves"] += 1
            content = _effective_content(node)
            resolved_image = False
            for content_index, item in enumerate(content):
                item_path = f"{path}.answer_content[{content_index}]"
                if isinstance(item, AnswerParagraph):
                    stats["paragraphs"] += 1
                elif isinstance(item, AnswerTable):
                    stats["tables"] += 1
                    if not item.rows or not any(item.rows):
                        issues.append(
                            {
                                "code": "empty_table",
                                "severity": "blocking",
                                "message": "An empty structured table would lose answer content.",
                                "path": item_path,
                            }
                        )
                elif isinstance(item, AnswerImage):
                    stats["images"] += 1
                    available = item.local_id in available_assets
                    resolved_image = resolved_image or available
                    if not available:
                        issues.append(
                            {
                                "code": "missing_asset",
                                "severity": "blocking",
                                "message": (
                                    f"Referenced image asset {item.local_id!r} is unavailable."
                                ),
                                "path": item_path,
                            }
                        )
                elif isinstance(item, UnsupportedAnswerContent):
                    issues.append(
                        {
                            "code": "unsupported_answer_content",
                            "severity": "blocking",
                            "message": f"Unsupported non-text answer content: {item.reason}.",
                            "path": item_path,
                        }
                    )
            if node.has_non_text_content and not resolved_image:
                issues.append(
                    {
                        "code": "unresolved_non_text",
                        "severity": "blocking",
                        "message": "Non-text answer content has no resolved image asset.",
                        "path": path,
                    }
                )
            for child_index, child in enumerate(node.children):
                walk(child, f"{path}.children[{child_index}]")

        for question_index, question in enumerate(section.questions):
            walk(question, f"sections[{section_index}].questions[{question_index}]")
    blocking = sum(issue["severity"] == "blocking" for issue in issues)
    return {"valid": blocking == 0, "blocking_count": blocking, "issues": issues, "stats": stats}


def _repeat_header(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _alignment(value: str) -> WD_ALIGN_PARAGRAPH:
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[value]


def _style_cell(
    cell: _Cell,
    *,
    alignment: str,
    font_size_pt: float,
    width_mm: float | None = None,
) -> None:
    if width_mm is not None:
        cell.width = Mm(width_mm)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.style = "QuestionBody"
        paragraph.alignment = _alignment(alignment)
        for run in paragraph.runs:
            run.font.size = Pt(font_size_pt)


def _render_answer_table(
    document: DocumentObject,
    content: AnswerTable,
    profile: RenderTemplateProfile,
) -> None:
    layout = profile.config.answer_sheet_layout_json
    columns = max(len(row) for row in content.rows)
    table = document.add_table(rows=len(content.rows), cols=columns)
    table.autofit = True
    if layout.answer_table_borders:
        set_table_borders(table)
    for row_index, values in enumerate(content.rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index] if column_index < len(values) else ""
            _style_cell(
                cell,
                alignment=layout.answer_table_alignment,
                font_size_pt=layout.answer_table_font_size_pt,
            )
    if layout.repeat_table_header and table.rows:
        _repeat_header(table.rows[0])


def _render_mcq(
    document: DocumentObject,
    items: list[tuple[str, str]],
    profile: RenderTemplateProfile,
) -> None:
    layout = profile.config.answer_sheet_layout_json
    columns = layout.mcq_columns
    row_count = math.ceil(len(items) / columns)
    table = document.add_table(rows=row_count + 1, cols=columns * 2)
    table.autofit = False
    if layout.mcq_borders:
        set_table_borders(table)
    for row in table.rows:
        row.height = Mm(layout.mcq_row_height_mm)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for column in range(columns):
        question = table.cell(0, column * 2)
        answer = table.cell(0, column * 2 + 1)
        question.text = "題號"
        answer.text = "答案"
        _style_cell(
            question,
            alignment=layout.mcq_alignment,
            font_size_pt=layout.mcq_font_size_pt,
            width_mm=layout.mcq_question_width_mm,
        )
        _style_cell(
            answer,
            alignment=layout.mcq_alignment,
            font_size_pt=layout.mcq_font_size_pt,
            width_mm=layout.mcq_answer_width_mm,
        )
    _repeat_header(table.rows[0])
    for row in range(row_count):
        for column in range(columns):
            index = row + column * row_count
            if index >= len(items):
                continue
            question = table.cell(row + 1, column * 2)
            answer = table.cell(row + 1, column * 2 + 1)
            question.text = items[index][0]
            answer.text = items[index][1]
            _style_cell(
                question,
                alignment=layout.mcq_alignment,
                font_size_pt=layout.mcq_font_size_pt,
                width_mm=layout.mcq_question_width_mm,
            )
            _style_cell(
                answer,
                alignment=layout.mcq_alignment,
                font_size_pt=layout.mcq_font_size_pt,
                width_mm=layout.mcq_answer_width_mm,
            )


def _render_node(
    document: DocumentObject,
    node: AnsNode,
    depth: int,
    profile: RenderTemplateProfile,
    answer_assets: AnswerAssetResolver,
) -> None:
    layout = profile.config.answer_sheet_layout_json
    style = "QuestionBody" if depth == 0 else "QuestionSubpart"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.left_indent = Mm(depth * layout.hierarchy_indent_mm)
    paragraph.add_run(node.label or "?")
    if node.children:
        total = node.total()
        if layout.show_parent_totals and total is not None:
            paragraph.add_run(
                "  "
                + _format_marks(
                    total, profile.config.question_style_config_json.marks_format
                )
            )
    else:
        marks = (
            _format_marks(node.marks, profile.config.question_style_config_json.marks_format)
            if node.marks is not None
            else "(? marks)"
        )
        if profile.config.question_style_config_json.marks_display == "inline":
            paragraph.add_run(f"  {marks}")
        else:
            marks_paragraph = document.add_paragraph(marks, style="QuestionMarks")
            if profile.config.question_style_config_json.marks_display == "below":
                marks_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for item in _effective_content(node):
        if isinstance(item, AnswerParagraph):
            answer = document.add_paragraph(item.text, style="QuestionBody")
            answer.paragraph_format.left_indent = Mm(
                (depth + 1) * layout.hierarchy_indent_mm
            )
        elif isinstance(item, AnswerTable):
            _render_answer_table(document, item, profile)
        elif isinstance(item, AnswerImage):
            image = document.add_paragraph(style="QuestionBody")
            image.paragraph_format.left_indent = Mm(
                (depth + 1) * layout.hierarchy_indent_mm
            )
            image.add_run().add_picture(
                BytesIO(answer_assets(item.local_id)), width=Mm(layout.image_max_width_mm)
            )
    for child in node.children:
        _render_node(document, child, depth + 1, profile, answer_assets)


def render_answer_sheet(
    sheet: AnsSheet,
    profile: RenderTemplateProfile,
    metadata: DocumentMetadata,
    answer_assets: Mapping[str, bytes],
    logo_assets: LogoAssetResolver,
) -> AnswerSheetRenderResult:
    validation = validate_answer_sheet_render(sheet, set(answer_assets))
    if not validation["valid"]:
        return AnswerSheetRenderResult(docx=None, validation=validation)

    document = Document()
    _configure_styles(document, profile.config)
    section = document.sections[0]
    page = profile.config.page_config_json
    if page.size == "A4":
        section.page_width, section.page_height = Mm(210), Mm(297)
    else:
        section.page_width, section.page_height = Mm(215.9), Mm(279.4)
    section.top_margin = Mm(page.margin_top_mm)
    section.right_margin = Mm(page.margin_right_mm)
    section.bottom_margin = Mm(page.margin_bottom_mm)
    section.left_margin = Mm(page.margin_left_mm)

    header = section.header.paragraphs[0]
    header.style = "PaperMetadata"
    if profile.logo_asset_id is not None:
        header.add_run().add_picture(BytesIO(logo_assets(profile.logo_asset_id)), width=Mm(20))
    header_text = _substitute(profile.config.header_config_json.text, metadata)
    if header_text:
        header.add_run(header_text)
    footer = section.footer.paragraphs[0]
    footer.style = "PaperMetadata"
    footer_text = _substitute(profile.config.footer_config_json.text, metadata)
    footer.add_run(footer_text)
    if profile.config.footer_config_json.page_numbering:
        if footer_text:
            footer.add_run("  ")
        append_page_number(footer)

    for line in profile.config.answer_sheet_layout_json.metadata_lines:
        value = _substitute(line, metadata)
        if value:
            document.add_paragraph(
                value,
                style="PaperTitle" if not document.paragraphs else "PaperMetadata",
            )

    for answer_section in sheet.sections:
        heading = answer_section.title
        computed = answer_section.computed_total()
        if (
            answer_section.declared_total is not None
            and computed is not None
            and answer_section.declared_total != computed
        ):
            heading += f"  [declared {answer_section.declared_total} / computed {computed}]"
        document.add_paragraph(heading, style="SectionHeading")
        if answer_section.mcq:
            _render_mcq(document, answer_section.mcq, profile)
        for question in answer_section.questions:
            _render_node(document, question, 0, profile, answer_assets.__getitem__)

    if sheet.warnings:
        document.add_paragraph("Validation warnings / 驗證警告", style="SectionHeading")
        for warning in sheet.warnings:
            document.add_paragraph(
                f"[{warning.code}] {warning.message}", style="QuestionBody"
            )

    output = BytesIO()
    document.save(output)
    return AnswerSheetRenderResult(
        docx=normalize_zip_timestamps(output.getvalue()), validation=validation
    )


__all__ = [
    "AnswerSheetRenderResult",
    "render_answer_sheet",
    "validate_answer_sheet_render",
]
