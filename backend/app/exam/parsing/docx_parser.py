"""DOCX → ``DocumentBlock[]`` parser.

Uses python-docx to walk the body in reading order (paragraphs and tables
interleaved), plus headers and footers. It performs layout detection only —
heading style, inline equation (OMML), inline image, caption, table — and never
interprets question structure or marks.

DOCX has no absolute page geometry without rendering, so every block's ``page``
and ``bbox`` are ``None`` (a documented "missing bbox" case); reading order is
preserved via the ``order`` index.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from app.exam.ir import AssetReference
from app.exam.parsing.base import ParseError
from app.exam.parsing.blocks import BlockKind, DocumentBlock, ParseResult, SourceReference

_HEADING_STYLE = re.compile(r"^Heading\s*(\d+)$", re.IGNORECASE)
_CAPTION_PREFIX = re.compile(
    r"^\s*(?:Figure|Fig\.|Table|Chart|圖|表|附圖)\b", re.IGNORECASE
)


def _paragraph_text(p_el: Any) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _paragraph_style_id(p_el: Any) -> str | None:
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is None:
        return None
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is None:
        return None
    val = p_style.get(qn("w:val"))
    return val if isinstance(val, str) else None


def _blips(element: Any) -> list[Any]:
    return list(element.findall(".//" + qn("a:blip")))


def _image_parts(element: Any, doc: Document) -> list[tuple[str, bytes, str]]:
    """Return (rid, blob, content_type) for each embedded image in ``element``."""
    found: list[tuple[str, bytes, str]] = []
    seen: set[str] = set()
    for blip in _blips(element):
        rid = blip.get(qn("r:embed"))
        if not rid or rid in seen:
            continue
        seen.add(rid)
        part = doc.part.related_parts.get(rid)
        if part is not None:
            found.append((str(rid), part.blob, str(part.content_type)))
    return found


def _non_text_kind(el: Any) -> str:
    """Detect image/diagram/drawing/object content in an element; ``""`` if none.

    Uses ``iter()`` (not ``find()``) because LibreOffice-generated drawings are
    wrapped in ``mc:AlternateContent`` which ``find()`` does not traverse.
    """
    for node in el.iter():
        if not isinstance(node.tag, str):
            continue
        local = node.tag.rsplit("}", 1)[-1]
        if local == "blip":
            return "image"
        if local in ("txbxContent", "txbx"):
            return "drawing"
        if local == "pict":
            return "image"
        if local == "object":
            return "object"
    return ""


def _cell_text(tc: Any) -> tuple[str, str]:
    """Return ``(text, non_text_kind)`` for a table cell.

    Preserves paragraph/line structure: multiple paragraphs are joined with
    ``\\n`` and ``w:br``/``w:cr`` become ``\\n``. A nested table is flattened to
    ``"cell | cell | …"`` rows joined with ``\\n``. Text inside a drawing/textbox
    (a diagram) is NOT extracted — the cell is flagged ``drawing`` instead of
    producing garbled text.
    """
    parts: list[str] = []
    kind = ""
    for child in tc:
        if child.tag == qn("w:p"):
            nk = _non_text_kind(child)
            if nk:
                kind = kind or nk
                continue
            buf: list[str] = []
            for node in child.iter():
                if node.tag in (qn("w:br"), qn("w:cr")):
                    buf.append("\n")
                elif node.tag == qn("w:tab"):
                    buf.append("\t")
                elif node.tag == qn("w:t"):
                    buf.append(node.text or "")
            text = "".join(buf).strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            for tr in child.findall(qn("w:tr")):
                row_cells: list[str] = []
                for nested_tc in tr.findall(qn("w:tc")):
                    nested_text, nested_kind = _cell_text(nested_tc)
                    if nested_text:
                        row_cells.append(nested_text)
                    if nested_kind:
                        kind = kind or nested_kind
                if row_cells:
                    parts.append(" | ".join(row_cells))
    return "\n".join(parts), kind


def _table_rows(tbl_el: Any) -> tuple[list[list[str]], dict[str, str]]:
    """Return ``(rows, non_text_cells)``; ``non_text_cells`` maps ``"row:col"``
    to a kind (``drawing``/``image``) for cells whose content is not text."""
    rows: list[list[str]] = []
    non_text: dict[str, str] = {}
    for r, tr in enumerate(tbl_el.findall(qn("w:tr"))):
        row: list[str] = []
        for c, tc in enumerate(tr.findall(qn("w:tc"))):
            text, kind = _cell_text(tc)
            row.append(text)
            if kind:
                non_text[f"{r}:{c}"] = kind
        rows.append(row)
    return rows, non_text


def _cell_content_blocks(tc: Any) -> list[dict[str, Any]]:
    """Lossless supported content inside an answer cell.

    Top-level paragraphs remain paragraphs and nested tables remain structured
    rows.  Images/drawings are represented separately by ``cell_assets`` and
    ``non_text_cells`` because their bytes are relationship-backed.
    """
    blocks: list[dict[str, Any]] = []
    for child in tc:
        if child.tag == qn("w:p"):
            if _non_text_kind(child):
                continue
            buf: list[str] = []
            for node in child.iter():
                if node.tag in (qn("w:br"), qn("w:cr")):
                    buf.append("\n")
                elif node.tag == qn("w:tab"):
                    buf.append("\t")
                elif node.tag == qn("w:t"):
                    buf.append(node.text or "")
            value = "".join(buf).strip()
            if value:
                blocks.append({"kind": "paragraph", "text": value})
        elif child.tag == qn("w:tbl"):
            rows, _ = _table_rows(child)
            blocks.append({"kind": "table", "rows": rows})
    return blocks


class DocxParser:
    name = "docx"

    def supports(self, content_type: str | None, filename: str | None) -> bool:
        name = (filename or "").lower()
        return name.endswith(".docx") or content_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def parse(self, data: bytes, source_name: str | None = None) -> ParseResult:
        if not data or data[:4] != b"PK\x03\x04":
            raise ParseError("not a DOCX (missing ZIP magic)")
        try:
            doc = Document(BytesIO(data))
        except Exception as exc:  # noqa: BLE001 - malformed input surfaces as ParseError
            raise ParseError(f"malformed DOCX: {exc}") from exc

        blocks: list[DocumentBlock] = []
        assets: dict[str, bytes] = {}
        order = 0
        paragraph_index = 0
        table_index = 0
        image_index = 0
        image_refs: dict[str, AssetReference] = {}

        def capture_images(element: Any, doc_obj: Document) -> list[tuple[str, AssetReference]]:
            nonlocal image_index
            captured: list[tuple[str, AssetReference]] = []
            for rid, blob, content_type in _image_parts(element, doc_obj):
                reference = image_refs.get(rid)
                if reference is None:
                    local_id = f"img-{image_index}"
                    image_index += 1
                    assets[local_id] = blob
                    reference = AssetReference(local_id=local_id, mime_type=content_type)
                    image_refs[rid] = reference
                captured.append((rid, reference))
            return captured

        def append_image_blocks(
            captured: list[tuple[str, AssetReference]], parent_table: str | None = None
        ) -> None:
            nonlocal order
            for rid, reference in captured:
                blocks.append(
                    DocumentBlock(
                        id=f"b{order:04d}",
                        kind=BlockKind.IMAGE,
                        order=order,
                        source=SourceReference(
                            file_name=source_name, element_id=f"img-{rid}"
                        ),
                        asset=reference,
                        meta={"parent_table": parent_table} if parent_table else {},
                    )
                )
                order += 1

        # Body in reading order (paragraphs and tables interleaved).
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph_index += 1
                text = _paragraph_text(child)
                element_id = f"p{paragraph_index}"
                source = SourceReference(file_name=source_name, element_id=element_id)
                if child.find(qn("m:oMath")) is not None:
                    math_text = "".join(
                        t.text or "" for t in child.iter(qn("m:t"))
                    )
                    blocks.append(
                        DocumentBlock(
                            id=f"b{order:04d}",
                            kind=BlockKind.EQUATION,
                            text=math_text,
                            order=order,
                            source=source,
                            meta={"raw": "omml"},
                        )
                    )
                    order += 1
                    continue
                if _blips(child):
                    append_image_blocks(capture_images(child, doc))
                    if text.strip():
                        blocks.append(
                            DocumentBlock(
                                id=f"b{order:04d}",
                                kind=BlockKind.TEXT,
                                text=text,
                                order=order,
                                source=source,
                            )
                        )
                        order += 1
                    continue
                style_id = _paragraph_style_id(child)
                heading = _HEADING_STYLE.match(style_id or "")
                if heading:
                    blocks.append(
                        DocumentBlock(
                            id=f"b{order:04d}",
                            kind=BlockKind.HEADING,
                            text=text,
                            order=order,
                            source=source,
                            meta={"level": int(heading.group(1))},
                        )
                    )
                    order += 1
                    continue
                if _CAPTION_PREFIX.match(text):
                    blocks.append(
                        DocumentBlock(
                            id=f"b{order:04d}",
                            kind=BlockKind.CAPTION,
                            text=text,
                            order=order,
                            source=source,
                        )
                    )
                    order += 1
                    continue
                if text.strip():
                    blocks.append(
                        DocumentBlock(
                            id=f"b{order:04d}",
                            kind=BlockKind.TEXT,
                            text=text,
                            order=order,
                            source=source,
                        )
                    )
                    order += 1
            elif child.tag == qn("w:tbl"):
                table_index += 1
                rows, non_text = _table_rows(child)
                meta: dict[str, Any] = {"non_text_cells": non_text} if non_text else {}
                cell_assets: dict[str, list[dict[str, str]]] = {}
                cell_content: dict[str, list[dict[str, Any]]] = {}
                table_images: list[tuple[str, AssetReference]] = []
                for row_index, row in enumerate(child.findall(qn("w:tr"))):
                    for column_index, cell in enumerate(row.findall(qn("w:tc"))):
                        key = f"{row_index}:{column_index}"
                        content = _cell_content_blocks(cell)
                        if content:
                            cell_content[key] = content
                        captured = capture_images(cell, doc)
                        if captured:
                            cell_assets[key] = [
                                reference.model_dump(mode="json")
                                for _, reference in captured
                            ]
                            table_images.extend(captured)
                if cell_assets:
                    meta["cell_assets"] = cell_assets
                if cell_content:
                    meta["cell_content"] = cell_content
                blocks.append(
                    DocumentBlock(
                        id=f"b{order:04d}",
                        kind=BlockKind.TABLE,
                        rows=rows,
                        order=order,
                        source=SourceReference(
                            file_name=source_name, element_id=f"tbl{table_index}"
                        ),
                        meta=meta,
                    )
                )
                order += 1
                append_image_blocks(table_images, f"tbl{table_index}")

        # Headers and footers as separate, order-continuing blocks.
        for section in doc.sections:
            header_index = 0
            for p in section.header.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                header_index += 1
                blocks.append(
                    DocumentBlock(
                        id=f"b{order:04d}",
                        kind=BlockKind.HEADER,
                        text=text,
                        order=order,
                        source=SourceReference(
                            file_name=source_name, element_id=f"hdr{header_index}"
                        ),
                    )
                )
                order += 1
            footer_index = 0
            for p in section.footer.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                footer_index += 1
                blocks.append(
                    DocumentBlock(
                        id=f"b{order:04d}",
                        kind=BlockKind.FOOTER,
                        text=text,
                        order=order,
                        source=SourceReference(
                            file_name=source_name, element_id=f"ftr{footer_index}"
                        ),
                    )
                )
                order += 1

        return ParseResult(
            blocks=blocks,
            assets=assets,
            page_count=None,
            source_name=source_name,
            format="docx",
            warnings=[],
            needs_review=False,
        )
