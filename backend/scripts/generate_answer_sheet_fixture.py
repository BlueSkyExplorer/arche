"""Generate a deterministic answer-sheet .docx fixture for E2E / regression tests.

Mirrors the real 余振強紀念中學 answer-key structure (20260515 table format) so the
E2E smoke exercises the same parser paths without needing LibreOffice:

- 甲部 多項選擇題 (30分): 題號/答案 grid, 30 items.
- 乙部 結構題 (50分): Q1..Q9 question tables with forward-filled
  題號/子題/子子題 labels, nested (b)(i)/(b)(ii), one structured table answer,
  and one inline image answer (Q9b) — declared 50 vs computed 51.

Run from `backend/`:

    uv run python scripts/generate_answer_sheet_fixture.py
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x04\x00\x00\x00\xb5\x1c\x0c\x02\x00\x00\x00\x0bIDATx\xda\x63\x64"
    b"\xf8\x0f\x00\x01\x05\x01\x01'\x18\xe3f\x00\x00\x00\x00IEND\xaeB`\x82"
)

FIXTURE = Path(__file__).resolve().parent.parent / "tests" / "fixtures"


def _set_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(11)


def _add_image(cell, png: bytes) -> None:
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(png))


def _add_mcq(doc: Document) -> None:
    doc.add_paragraph("甲部 多項選擇題 (30分)")
    table = doc.add_table(rows=16, cols=4)
    for idx, header in enumerate(["題號", "答案", "題號", "答案"]):
        _set_text(table.cell(0, idx), header)
    answers = ["A", "B", "C", "D"]
    for row in range(1, 16):
        left, right = row, row + 15
        _set_text(table.cell(row, 0), str(left))
        _set_text(table.cell(row, 1), answers[(left - 1) % 4])
        _set_text(table.cell(row, 2), str(right))
        _set_text(table.cell(row, 3), answers[(right - 1) % 4])


def _add_nested_table(cell) -> None:
    """Append a native nested w:tbl (構造|風媒花|蟲媒花) into ``cell``."""
    helper = Document()
    nested = helper.add_table(rows=2, cols=3)
    for row in range(2):
        for col in range(3):
            nested.cell(row, col).text = ""
    nested.cell(0, 0).text = "構造"
    nested.cell(0, 1).text = "風媒花"
    nested.cell(0, 2).text = "蟲媒花"
    nested.cell(1, 0).text = "花瓣"
    nested.cell(1, 1).text = "細小"
    nested.cell(1, 2).text = "鮮艷"
    cell._tc.append(nested._tbl)  # type: ignore[attr-defined]
    cell._tc.append(OxmlElement("w:p"))  # type: ignore[attr-defined]


def _question_table(
    doc: Document,
    label: str,
    rows: list[tuple[str, str, str]],
    marks: dict[str, str],
) -> None:
    """One question table with forward-filled label columns (no header row).

    The deterministic extractor's ``_is_question_table`` expects the FIRST row's
    first cell to be the question label (``Q1``…), matching the real 20260515
    structure; a 題號/… header row would defeat detection.
    """
    table = doc.add_table(rows=len(rows), cols=5)
    for row_index, (sub, subsub, answer) in enumerate(rows):
        _set_text(table.cell(row_index, 0), label if row_index == 0 else "")
        _set_text(table.cell(row_index, 1), sub)
        _set_text(table.cell(row_index, 2), subsub)
        _set_text(table.cell(row_index, 3), answer)
        leaf = subsub or sub
        _set_text(table.cell(row_index, 4), marks.get(leaf, ""))


def build() -> bytes:
    doc = Document()
    _add_mcq(doc)

    doc.add_paragraph("")
    doc.add_paragraph("乙部 結構題 (50分)")

    # Q1 = 6 marks, nested (a) + (b)(i)/(b)(ii)
    _question_table(
        doc,
        "Q1",
        [
            ("(a)", "", "風媒花的適應特徵"),
            ("(b)", "(i)", "花瓣細小"),
            ("(b)", "(ii)", "柱頭呈羽狀"),
        ],
        {"(a)": "2分", "(i)": "2分", "(ii)": "2分"},
    )

    # Q2 = 5 marks, structured table answer (構造|風媒花|蟲媒花)
    q2 = doc.add_table(rows=1, cols=5)
    _set_text(q2.cell(0, 0), "Q2")
    _set_text(q2.cell(0, 1), "(a)")
    _set_text(q2.cell(0, 2), "")
    _add_nested_table(q2.cell(0, 3))
    _set_text(q2.cell(0, 4), "5分")

    # Q3..Q8 single-leaf questions: 3, 9, 7, 5, 9, 3
    for label, mark in [("Q3", 3), ("Q4", 9), ("Q5", 7), ("Q6", 5), ("Q7", 9), ("Q8", 3)]:
        _question_table(doc, label, [("", "", f"{label} 答案")], {"": f"{mark}分"})

    # Q9 = 4 marks: (a) text + (b) inline image answer
    q9 = doc.add_table(rows=2, cols=5)
    _set_text(q9.cell(0, 0), "Q9")
    _set_text(q9.cell(0, 1), "(a)")
    _set_text(q9.cell(0, 2), "")
    _set_text(q9.cell(0, 3), "染色體數目")
    _set_text(q9.cell(0, 4), "2分")
    _set_text(q9.cell(1, 1), "(b)")
    _set_text(q9.cell(1, 2), "")
    _add_image(q9.cell(1, 3), PNG)
    _set_text(q9.cell(1, 4), "2分")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def main() -> None:
    from app.document.ooxml.package import normalize_zip_timestamps

    FIXTURE.mkdir(parents=True, exist_ok=True)
    target = FIXTURE / "answer_sheet.docx"
    target.write_bytes(normalize_zip_timestamps(build()))
    print(f"wrote {target} ({target.stat().st_size} bytes)")


if __name__ == "__main__":
    main()