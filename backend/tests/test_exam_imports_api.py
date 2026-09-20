"""Exam-import workflow API tests (upload -> review -> approve -> Question Library)."""

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models import Question


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("1. State one function of the cell membrane. (5 marks)")
    doc.add_paragraph("(a) Define diffusion. (2 marks)")
    doc.add_paragraph("(b) Define osmosis. (3 marks)")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(api_client: TestClient, data: bytes, filename: str = "biology.docx"):
    return api_client.post(
        "/api/v1/exam-imports",
        files={"file": (filename, data, "application/octet-stream")},
    )


@pytest.mark.db
def test_import_workflow_happy_path(
    api_client: TestClient, db_session: Session
) -> None:
    # 1. upload -> parse -> extract -> validate
    response = _upload(api_client, _docx_bytes())
    assert response.status_code in (200, 201), response.text
    imp = response.json()
    assert imp["status"] in ("needs_review", "ready")
    assert imp["extractor_name"] == "rule-based"  # AI disabled -> rule-based
    assert imp["blocks_json"]  # DocumentBlock evidence persisted
    assert imp["reviewed_document_json"] == imp["extracted_document_json"]

    # 2. list + detail
    listing = api_client.get("/api/v1/exam-imports").json()
    assert any(item["id"] == imp["id"] for item in listing)
    detail = api_client.get(f"/api/v1/exam-imports/{imp['id']}").json()
    assert detail["reviewed_document_json"] is not None

    # 3. review: edit the reviewed document (change sub-question (a) label -> (i))
    reviewed = detail["reviewed_document_json"]
    for section in reviewed["sections"]:
        for question in section["questions"]:
            for child in question["children"]:
                if child["label"] == "(a)":
                    child["label"] = "(i)"
    put = api_client.put(
        f"/api/v1/exam-imports/{imp['id']}/reviewed", json=reviewed
    )
    assert put.status_code == 200, put.text
    assert put.json()["reviewed_document_json"] != imp["extracted_document_json"]

    # 4. reload -> edit persisted
    reloaded = api_client.get(f"/api/v1/exam-imports/{imp['id']}").json()
    labels = [
        c["label"]
        for s in reloaded["reviewed_document_json"]["sections"]
        for q in s["questions"]
        for c in q["children"]
    ]
    assert "(i)" in labels and "(a)" not in labels

    # 5. approve -> materialize -> persist questions
    approved = api_client.post(f"/api/v1/exam-imports/{imp['id']}/approve")
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "completed"
    assert len(approved.json()["created_question_ids"]) == 1

    # 6. Question Library has the created question with correct leaf marks
    qid = approved.json()["created_question_ids"][0]
    question = db_session.get(Question, qid)
    assert question is not None
    content = question.content_json
    sub_marks = {
        block["attrs"]["label"]: block["attrs"].get("marks")
        for block in content["content"]
        if block.get("type") == "subQuestion"
    }
    assert sub_marks == {"(i)": "2", "(b)": "3"}
    assert content.get("marks") is None  # non-leaf: computed from leaves


@pytest.mark.db
def test_import_unknown_marks_stay_null(api_client: TestClient, db_session: Session) -> None:
    doc = Document()
    doc.add_paragraph("1. A question with no marks.")
    buf = BytesIO()
    doc.save(buf)
    response = _upload(api_client, buf.getvalue())
    assert response.status_code in (200, 201), response.text
    imp = response.json()
    assert imp["needs_review"] is True
    # the extracted doc's leaf own_marks is null (not 0)
    leaves = [
        q
        for s in imp["reviewed_document_json"]["sections"]
        for q in s["questions"]
    ]
    assert leaves[0]["own_marks"] is None


@pytest.mark.db
def test_approve_blocked_by_missing_asset(api_client: TestClient) -> None:
    response = _upload(api_client, _docx_bytes())
    imp = response.json()
    # inject an image whose local_id is NOT in the persisted asset manifest
    reviewed = imp["reviewed_document_json"]
    reviewed["sections"][0]["questions"][0]["content"].append(
        {"kind": "image", "asset": {"local_id": "ghost-img", "mime_type": "image/png"}}
    )
    reviewed.setdefault("assets", []).append(
        {"local_id": "ghost-img", "mime_type": "image/png"}
    )
    put = api_client.put(f"/api/v1/exam-imports/{imp['id']}/reviewed", json=reviewed)
    assert put.status_code == 200, put.text
    approve = api_client.post(f"/api/v1/exam-imports/{imp['id']}/approve")
    assert approve.status_code == 409
    assert "asset" in approve.json()["detail"]
