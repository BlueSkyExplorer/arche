from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.storage import LocalDirStorage
from app.models import ExamImport, Workspace
from tests.conftest import switch_workspace

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x04\x00\x00\x00\xb5\x1c\x0c\x02\x00\x00\x00\x0bIDATx\xda\x63\x64"
    b"\xf8\x0f\x00\x01\x05\x01\x01'\x18\xe3f\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _sheet(answer: str = "REVIEWED ANSWER", image_id: str = "img-0") -> dict[str, Any]:
    questions = []
    totals = [6, 5, 3, 9, 7, 5, 9, 3, 4]
    for index, marks in enumerate(totals, 1):
        content: list[dict[str, Any]] = [
            {"kind": "paragraph", "text": answer if index == 1 else f"answer {index}"}
        ]
        if index == 1:
            content.append(
                {
                    "kind": "table",
                    "rows": [["構造", "風媒花", "蟲媒花"], ["花瓣", "細小", "鮮艷"]],
                }
            )
        if index == 9:
            content.append({"kind": "image", "local_id": image_id, "mime_type": "image/png"})
        if index == 1:
            questions.append(
                {
                    "label": "Q1",
                    "answer": [],
                    "answer_content": [],
                    "marks": None,
                    "children": [
                        {
                            "label": "(a)",
                            "answer": [answer],
                            "answer_content": content,
                            "marks": "2",
                            "children": [],
                            "has_non_text_content": False,
                        },
                        {
                            "label": "(b)",
                            "answer": [],
                            "answer_content": [],
                            "marks": None,
                            "children": [
                                {
                                    "label": "(i)",
                                    "answer": ["nested answer"],
                                    "answer_content": [
                                        {"kind": "paragraph", "text": "nested answer"}
                                    ],
                                    "marks": "4",
                                    "children": [],
                                    "has_non_text_content": False,
                                }
                            ],
                            "has_non_text_content": False,
                        },
                    ],
                    "has_non_text_content": False,
                }
            )
        else:
            questions.append(
                {
                    "label": f"Q{index}",
                    "answer": [f"answer {index}"],
                    "answer_content": content,
                    "marks": str(marks),
                    "children": [],
                    "has_non_text_content": index == 9,
                }
            )
    return {
        "title": "Source title",
        "sections": [
            {
                "title": "甲部 多項選擇題 (30分)",
                "declared_total": "30",
                "computed_total": None,
                "mcq": [[str(index), "ABCD"[(index - 1) % 4]] for index in range(1, 31)],
                "questions": [],
                "standalone_non_text": False,
            },
            {
                "title": "乙部 結構題 (50分)",
                "declared_total": "50",
                "computed_total": "51",
                "mcq": None,
                "questions": questions,
                "standalone_non_text": False,
            },
        ],
        "warnings": [
            {
                "code": "declared_total_mismatch",
                "message": "declares 50 marks but extracted leaf marks sum to 51",
                "section": "乙部 結構題 (50分)",
                "path": None,
                "declared": "50",
                "computed": "51",
            }
        ],
        "asset_refs": [{"local_id": image_id, "mime_type": "image/png"}],
    }


def _completed_import(
    client: TestClient, db: Session, *, missing_asset: bool = False
) -> ExamImport:
    storage = LocalDirStorage(client.test_settings.storage_local_dir)  # type: ignore[attr-defined]
    key = f"exam-imports/{client.workspace_id}/fixture/assets/img-0"  # type: ignore[attr-defined]
    if not missing_asset:
        storage.put(key, PNG)
    item = ExamImport(
        workspace_id=client.workspace_id,  # type: ignore[attr-defined]
        source_filename="20260515 (C) ANS.doc",
        source_type="docx",
        import_type="answer_sheet",
        status="completed",
        answer_sheet_json=_sheet("ORIGINAL EXTRACTED"),
        reviewed_answer_sheet_json=_sheet(),
        asset_manifest={"img-0": {"storage_key": key, "mime_type": "image/png"}},
        completed_at=datetime.now(UTC),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def _metadata(template_id: str) -> dict[str, Any]:
    return {
        "template_profile_id": template_id,
        "metadata": {
            "school_name": "余振強紀念中學",
            "academic_year": "2024-2025",
            "exam_name": "下學期考試",
            "level": "中四級",
            "subject": "生物科",
            "document_type": "參考答案",
        },
    }


@pytest.mark.db
def test_completed_answer_sheet_preview_and_export_are_lossless(
    api_client: TestClient, db_session: Session, template_payload: dict[str, Any]
) -> None:
    template_payload["answer_sheet_layout_json"] = {
        "metadata_lines": [
            "{{school_name}}",
            "{{academic_year}} {{exam_name}}",
            "{{level}} {{subject}} ({{document_type}})",
        ],
        "mcq_columns": 2,
        "mcq_borders": True,
        "mcq_question_width_mm": 13,
        "mcq_answer_width_mm": 17,
        "mcq_alignment": "center",
        "mcq_font_size_pt": 10,
        "mcq_row_height_mm": 6,
        "hierarchy_indent_mm": 7,
        "show_parent_totals": True,
        "image_max_width_mm": 100,
        "answer_table_borders": True,
        "answer_table_alignment": "left",
        "answer_table_font_size_pt": 9,
        "repeat_table_header": True,
    }
    template = api_client.post("/api/v1/templates", json=template_payload).json()
    item = _completed_import(api_client, db_session)
    asset = api_client.get(f"/api/v1/exam-imports/{item.id}/assets/img-0")
    assert asset.status_code == 200
    assert asset.headers["content-type"] == "image/png"
    preview = api_client.post(
        f"/api/v1/exam-imports/{item.id}/answer-sheet-preview",
        json=_metadata(template["id"]),
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["record"]["status"] == "succeeded"
    assert body["record"]["validation_json"]["valid"] is True
    assert body["record"]["validation_json"]["stats"] == {
        "sections": 2,
        "mcq_items": 30,
        "question_nodes": 12,
        "leaves": 10,
        "paragraphs": 10,
        "tables": 1,
        "images": 1,
    }
    assert body["preview"]["answer_sheet"]["sections"][1]["declared_total"] == "50"
    assert body["preview"]["answer_sheet"]["sections"][1]["computed_total"] == "51"

    exported = api_client.post(
        f"/api/v1/exam-imports/{item.id}/answer-sheet-exports",
        json=_metadata(template["id"]),
    ).json()
    download = api_client.get(f"/api/v1/answer-sheet-exports/{exported['id']}/download")
    assert download.status_code == 200
    assert download.content.startswith(b"PK")
    document = Document(io.BytesIO(download.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "REVIEWED ANSWER" in text
    assert "nested answer" in text
    assert "ORIGINAL EXTRACTED" not in text
    # formal DOCX stays source-facing: no audit annotation or warning appendix
    assert "乙部 結構題 (50分)" in text
    assert "declared 50 / computed 51" not in text
    assert "Validation warnings" not in text
    assert "驗證警告" not in text
    # …but the mismatch fingerprint is still preserved on the export record
    codes = [issue["code"] for issue in exported["validation_json"]["issues"]]
    assert "declared_total_mismatch" in codes
    assert "構造" in table_text and "風媒花" in table_text and "蟲媒花" in table_text
    assert sum(value in table_text for value in ("1", "15", "16", "30")) == 4
    with zipfile.ZipFile(io.BytesIO(download.content)) as package:
        assert package.testzip() is None
        assert any(name.startswith("word/media/") for name in package.namelist())
        assert b"relationships/image" in package.read("word/_rels/document.xml.rels")


@pytest.mark.db
def test_missing_answer_asset_blocks_export(
    api_client: TestClient, db_session: Session, template_payload: dict[str, Any]
) -> None:
    template = api_client.post("/api/v1/templates", json=template_payload).json()
    item = _completed_import(api_client, db_session, missing_asset=True)
    response = api_client.post(
        f"/api/v1/exam-imports/{item.id}/answer-sheet-preview",
        json=_metadata(template["id"]),
    )
    assert response.status_code == 200
    record = response.json()["record"]
    assert record["status"] == "blocked"
    assert record["validation_json"]["valid"] is False
    assert any(
        issue["code"] == "missing_asset"
        for issue in record["validation_json"]["issues"]
    )
    download = api_client.get(f"/api/v1/answer-sheet-exports/{record['id']}/download")
    assert download.status_code == 409


@pytest.mark.db
def test_template_snapshot_and_workspace_isolation(
    api_client: TestClient, db_session: Session, template_payload: dict[str, Any]
) -> None:
    template = api_client.post("/api/v1/templates", json=template_payload).json()
    item = _completed_import(api_client, db_session)
    first = api_client.post(
        f"/api/v1/exam-imports/{item.id}/answer-sheet-preview",
        json=_metadata(template["id"]),
    ).json()["record"]
    patched = api_client.patch(
        f"/api/v1/templates/{template['id']}",
        json={"footer_config_json": {"text": "Changed later", "page_numbering": False}},
    )
    assert patched.json()["version"] == template["version"] + 1
    reloaded = api_client.get(f"/api/v1/answer-sheet-exports/{first['id']}").json()
    assert reloaded["template_version"] == template["version"]
    footer = reloaded["template_config_snapshot"]["config"]["footer_config_json"]
    assert footer["text"] == "Test footer"

    other = Workspace(name="Other", owner_user_id="other@example.com")
    db_session.add(other)
    db_session.commit()
    switch_workspace(api_client, other.id)
    assert api_client.get(f"/api/v1/exam-imports/{item.id}").status_code == 404
    assert api_client.get(f"/api/v1/answer-sheet-exports/{first['id']}").status_code == 404
    assert api_client.get(f"/api/v1/answer-sheet-exports/{first['id']}/download").status_code == 404
    assert api_client.get(f"/api/v1/exam-imports/{item.id}/assets/img-0").status_code == 404


@pytest.mark.db
def test_completed_answer_sheet_review_is_read_only(
    api_client: TestClient, db_session: Session
) -> None:
    item = _completed_import(api_client, db_session)
    response = api_client.put(
        f"/api/v1/exam-imports/{item.id}/reviewed", json=item.reviewed_answer_sheet_json
    )
    assert response.status_code == 409


@pytest.mark.db
def test_warning_counts_and_invalid_import_type(
    api_client: TestClient, db_session: Session
) -> None:
    item = _completed_import(api_client, db_session)
    item.reviewed_answer_sheet_json = {**_sheet(), "warnings": []}
    db_session.commit()
    detail = api_client.get(f"/api/v1/exam-imports/{item.id}").json()
    assert detail["original_warning_count"] == 1
    assert detail["current_warning_count"] == 0
    assert detail["warning_count"] == 0

    response = api_client.post(
        "/api/v1/exam-imports",
        data={"import_type": "garbage"},
        files={"file": ("paper.docx", b"PK-invalid", "application/octet-stream")},
    )
    assert response.status_code == 422


def _render_config(header_text: str, footer_text: str, metadata_lines: list[str]):
    from app.schemas.template_profile import TemplateProfileConfig

    return TemplateProfileConfig.model_validate(
        {
            "page_config_json": {
                "size": "A4", "margin_top_mm": 20, "margin_right_mm": 20,
                "margin_bottom_mm": 20, "margin_left_mm": 20,
            },
            "typography_config_json": {
                "chinese_font": "Noto Sans CJK TC", "latin_font": "Arial",
                "base_font_size_pt": 11, "line_spacing": 1.15,
            },
            "header_config_json": {"text": header_text},
            "footer_config_json": {"text": footer_text, "page_numbering": False},
            "numbering_config_json": {
                "question_style": "1.", "sub_question_style": "(a)",
                "sub_sub_question_style": "roman",
            },
            "section_style_config_json": {"spacing_before_pt": 6, "spacing_after_pt": 6},
            "question_style_config_json": {
                "spacing_before_pt": 3, "spacing_after_pt": 3,
                "marks_display": "right", "marks_format": "({marks} marks)",
                "default_answer_lines": 0,
            },
            "answer_sheet_layout_json": {"metadata_lines": metadata_lines},
            "role_styles": {},
        }
    )


def test_answer_sheet_render_is_xml_safe_for_metadata_special_chars() -> None:
    """`&`/`<`/`>` in user-supplied metadata/substitutions must round-trip and
    stay well-formed; unknown placeholders must stay literal, never error."""
    import xml.etree.ElementTree as ET

    from app.document.renderer import RenderTemplateProfile
    from app.document.renderer.answer_sheet import render_answer_sheet
    from app.exam.extraction.answer_sheet import sheet_from_dict
    from app.schemas.answer_sheet_export import DocumentMetadata

    sheet = sheet_from_dict(
        {
            "title": "t",
            "sections": [
                {
                    "title": "乙部 結構題 (50分)", "declared_total": "50", "mcq": None,
                    "questions": [
                        {
                            "label": "Q1", "answer": [], "answer_content": [], "marks": None,
                            "children": [
                                {
                                    "label": "(a)", "answer": ["A&B <C>"], "marks": "6",
                                    "answer_content": [{"kind": "paragraph", "text": "A&B <C>"}],
                                    "children": [], "has_non_text_content": False,
                                },
                            ],
                            "has_non_text_content": False,
                        }
                    ],
                    "standalone_non_text": False,
                }
            ],
            "warnings": [], "asset_refs": [],
        }
    )
    config = _render_config(
        header_text="期中考 & <{{subject}}>",
        footer_text="footer {{unknown_placeholder}}",
        metadata_lines=["{{school_name}} & {{subject}}"],
    )
    metadata = DocumentMetadata(
        school_name="A&B <校>", academic_year="2024-25", exam_name="Final",
        level="S4", subject="Bio <&>", document_type="Answer",
    )
    profile = RenderTemplateProfile(school_name="", logo_asset_id=None, config=config)
    result = render_answer_sheet(sheet, profile, metadata, {}, lambda _: b"")
    assert result.docx is not None, result.validation

    document = Document(io.BytesIO(result.docx))
    body_text = "\n".join(p.text for p in document.paragraphs)
    header_text = "\n".join(p.text for p in document.sections[0].header.paragraphs)
    footer_text = "\n".join(p.text for p in document.sections[0].footer.paragraphs)

    assert "A&B <C>" in body_text  # answer text round-trips
    assert "A&B <校>" in body_text  # substituted metadata line (school_name)
    assert "Bio <&>" in header_text  # {{subject}} substituted into header
    assert "{{unknown_placeholder}}" in footer_text  # unknown placeholder stays literal
    with zipfile.ZipFile(io.BytesIO(result.docx)) as package:
        ET.fromstring(package.read("word/document.xml"))  # well-formed XML
