from __future__ import annotations

import os
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from app.document.renderer import (
    PaperData,
    RenderQuestion,
    RenderSection,
    _format_marks,
    render_paper,
)
from app.schemas.content import DocNode
from app.schemas.template_profile import NumberingConfig as TemplateNumberingConfig
from app.services.numbering import (
    NumberingConfig,
    QuestionForNumbering,
    SectionForNumbering,
    number_questions,
)
from tests.document.fixture import FIXTURES, IMAGE_ID, paper, profile
from tests.document.regenerate import update_golden

GOLDENS = Path(__file__).parent / "goldens"


def _part(docx: bytes, name: str) -> bytes:
    with ZipFile(BytesIO(docx)) as archive:
        return archive.read(name)


def _assert_golden(name: str, actual: bytes) -> None:
    path = GOLDENS / name
    if os.getenv("REGEN_GOLDEN") == "1":
        update_golden(path, actual)
    assert actual == path.read_bytes()


def _render_content(content: DocNode) -> bytes:
    question = RenderQuestion(
        id=paper().sections[0].questions[0].id,
        position=1,
        content=content,
        marks=Decimal("1"),
    )
    data = PaperData(
        title="Renderer edge case",
        sections=(RenderSection(title="Section", position=1, questions=(question,)),),
    )
    return render_paper(data, profile(), {IMAGE_ID: (FIXTURES / "tiny.png").read_bytes()})


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1", "(1 marks)"),
        ("2", "(2 marks)"),
        ("7", "(7 marks)"),
        ("10", "(10 marks)"),
        ("20", "(20 marks)"),
        ("100", "(100 marks)"),
        ("0.5", "(0.5 marks)"),
        ("2.50", "(2.5 marks)"),
        ("12.25", "(12.25 marks)"),
    ],
)
def test_format_marks_preserves_integer_zeros(value: str, expected: str) -> None:
    assert _format_marks(Decimal(value), "({marks} marks)") == expected


def test_default_profile_golden_and_wording_fidelity() -> None:
    output = render_paper(paper(), profile(), {IMAGE_ID: (FIXTURES / "tiny.png").read_bytes()})
    _assert_golden("default-document.xml", _part(output, "word/document.xml"))
    extracted = "\n".join(paragraph.text for paragraph in Document(BytesIO(output)).paragraphs)
    for wording in (
        "保留原文 Keep wording: ",
        "粗體",
        " italic",
        " underline",
        "第一小題",
        "Second part",
        "第三小題",
        "Ten-mark question",
        "(10 marks)",
        "Twenty-mark question",
        "(20 marks)",
    ):
        assert wording in extracted


def test_variant_profile_golden_proves_page_and_typography_config() -> None:
    output = render_paper(
        paper(), profile(variant=True), {IMAGE_ID: (FIXTURES / "tiny.png").read_bytes()}
    )
    _assert_golden("variant-document.xml", _part(output, "word/document.xml"))
    styles = _part(output, "word/styles.xml")
    assert b"Times New Roman" in styles
    assert b"PMingLiU" in styles


def test_render_is_byte_for_byte_deterministic() -> None:
    assets = {IMAGE_ID: (FIXTURES / "tiny.png").read_bytes()}
    first = render_paper(paper(), profile(), assets)
    second = render_paper(paper(), profile(), assets)
    assert _part(first, "word/document.xml") == _part(second, "word/document.xml")
    assert first == second


def test_camel_case_image_asset_is_embedded_in_document() -> None:
    image = (FIXTURES / "tiny.png").read_bytes()
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {"type": "image", "attrs": {"assetId": str(IMAGE_ID), "alt": None, "widthMm": 20}}
            ],
        }
    )
    rendered = _render_content(content)
    with ZipFile(BytesIO(rendered)) as archive:
        assert image in (
            archive.read(name) for name in archive.namelist() if name.startswith("word/media/")
        )
        assert b"r:embed" in archive.read("word/document.xml")


def test_blank_height_answer_space_is_rendered() -> None:
    rendered = _render_content(
        DocNode.model_validate(
            {
                "type": "doc",
                "content": [{"type": "answerSpace", "attrs": {"blankHeightMm": 12.5}}],
            }
        )
    )
    assert b'w:line="709"' in _part(rendered, "word/document.xml")


def test_tiptap_link_mark_renders_text_and_hyperlink_relationship() -> None:
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "paragraph",
                    "content": [
                        {
                            "type": "text",
                            "text": "linked text",
                            "marks": [
                                {
                                    "type": "link",
                                    "attrs": {
                                        "href": "https://example.com",
                                        "target": "_blank",
                                        "rel": "noopener noreferrer",
                                    },
                                }
                            ],
                        }
                    ],
                }
            ],
        }
    )
    rendered = _render_content(content)
    assert b"linked text" in _part(rendered, "word/document.xml")
    rels = _part(rendered, "word/_rels/document.xml.rels")
    assert b'Target="https://example.com"' in rels
    assert b'rel="noopener noreferrer"' not in _part(rendered, "word/document.xml")


def test_ooxml_border_elements_follow_schema_child_order() -> None:
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {"type": "answerSpace", "attrs": {"lines": 1}},
                {
                    "type": "table",
                    "content": [
                        {
                            "type": "tableRow",
                            "content": [{"type": "tableCell", "content": [{"type": "paragraph"}]}],
                        }
                    ],
                },
            ],
        }
    )
    xml = _part(_render_content(content), "word/document.xml")
    answer_ppr = xml[xml.index(b'<w:pPr><w:pStyle w:val="AnswerSpace"') :]
    answer_ppr = answer_ppr[: answer_ppr.index(b"</w:pPr>")]
    assert answer_ppr.index(b"<w:pBdr>") < answer_ppr.index(b"<w:spacing")
    table_pr = xml[xml.index(b"<w:tblPr>") : xml.index(b"</w:tblPr>")]
    assert table_pr.index(b"<w:tblBorders>") < table_pr.index(b"<w:tblLook")


@pytest.mark.parametrize(
    "ending",
    [
        {
            "type": "table",
            "content": [
                {
                    "type": "tableRow",
                    "content": [{"type": "tableCell", "content": [{"type": "paragraph"}]}],
                }
            ],
        },
        {"type": "answerSpace", "attrs": {"lines": 1}},
    ],
)
def test_inline_marks_after_non_paragraph_final_block(ending: object) -> None:
    config = profile().config.model_copy(
        update={
            "question_style_config_json": profile().config.question_style_config_json.model_copy(
                update={"marks_display": "inline"}
            )
        }
    )
    rendered = render_paper(
        PaperData(
            title="Marks placement",
            sections=(
                RenderSection(
                    title="Section",
                    position=1,
                    questions=(
                        RenderQuestion(
                            id=paper().sections[0].questions[0].id,
                            position=1,
                            content=DocNode.model_validate(
                                {
                                    "type": "doc",
                                    "content": [
                                        {
                                            "type": "paragraph",
                                            "content": [{"type": "text", "text": "stem"}],
                                        },
                                        ending,
                                    ],
                                }
                            ),
                            marks=Decimal("1"),
                        ),
                    ),
                ),
            ),
        ),
        profile().model_copy(update={"config": config}),
        {},
    )
    xml = _part(rendered, "word/document.xml")
    final_block = (
        b"<w:tbl>"
        if isinstance(ending, dict) and ending["type"] == "table"
        else b'w:pStyle w:val="AnswerSpace"'
    )
    marks = b"(1 marks)"
    assert xml.index(final_block) < xml.index(marks)
    marks_context = xml[xml.rindex(b"<w:p", 0, xml.index(marks)) : xml.index(marks)]
    assert b"QuestionMarks" in marks_context


def test_default_answer_lines_are_appended_only_when_content_has_none() -> None:
    config = profile().config.model_copy(
        update={
            "question_style_config_json": profile().config.question_style_config_json.model_copy(
                update={"default_answer_lines": 2}
            )
        }
    )
    questions = (
        RenderQuestion(
            id=paper().sections[0].questions[0].id,
            position=1,
            content=DocNode.model_validate({"type": "doc", "content": [{"type": "paragraph"}]}),
            marks=Decimal(1),
        ),
        RenderQuestion(
            id=paper().sections[0].questions[1].id,
            position=2,
            content=DocNode.model_validate(
                {
                    "type": "doc",
                    "content": [{"type": "answerSpace", "attrs": {"lines": 1}}],
                }
            ),
            marks=Decimal(1),
        ),
    )
    rendered = render_paper(
        PaperData(
            title="Defaults",
            sections=(RenderSection(title="S", position=1, questions=questions),),
        ),
        profile().model_copy(update={"config": config}),
        {},
    )
    assert _part(rendered, "word/document.xml").count(b'w:pStyle w:val="AnswerSpace"') == 3


def test_unlabelled_subquestions_use_configured_style_and_explicit_labels_win() -> None:
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "subQuestion",
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "auto one"}]}
                    ],
                },
                {
                    "type": "subQuestion",
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "auto two"}]}
                    ],
                },
                {
                    "type": "subQuestion",
                    "attrs": {"label": "Custom"},
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "explicit"}]}
                    ],
                },
            ],
        }
    )
    config = profile().config.model_copy(
        update={
            "numbering_config_json": TemplateNumberingConfig(
                question_style="1.", sub_question_style="upper-alpha"
            )
        }
    )
    rendered = render_paper(
        PaperData(
            title="Sub labels",
            sections=(
                RenderSection(
                    title="S",
                    position=1,
                    questions=(
                        RenderQuestion(
                            id=paper().sections[0].questions[0].id,
                            position=1,
                            content=content,
                            marks=Decimal(1),
                        ),
                    ),
                ),
            ),
        ),
        profile().model_copy(update={"config": config}),
        {},
    )
    text = "\n".join(p.text for p in Document(BytesIO(rendered)).paragraphs)
    assert "(A) auto one" in text
    assert "(B) auto two" in text
    assert "Custom explicit" in text


def test_nested_sub_questions_use_depth2_numbering() -> None:
    """Nested sub-questions use depth-2 counter, not resetting depth-1."""
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "subQuestion",
                    "attrs": {"label": "(a)"},
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "part a stem"}]},
                        {
                            "type": "subQuestion",
                            "attrs": {"label": "(i)"},
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": "nested i"}],
                                }
                            ],
                        },
                        {
                            "type": "subQuestion",
                            "attrs": {"label": "(ii)"},
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": "nested ii"}],
                                }
                            ],
                        },
                    ],
                },
                {
                    "type": "subQuestion",
                    "attrs": {"label": "(b)"},
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "part b stem"}]},
                    ],
                },
            ],
        }
    )
    config = profile().config.model_copy(
        update={
            "numbering_config_json": TemplateNumberingConfig(
                question_style="1.", sub_question_style="(a)", sub_sub_question_style="roman"
            )
        }
    )
    rendered = render_paper(
        PaperData(
            title="Nested",
            sections=(
                RenderSection(
                    title="S",
                    position=1,
                    questions=(
                        RenderQuestion(
                            id=paper().sections[0].questions[0].id,
                            position=1,
                            content=content,
                            marks=Decimal(2),
                        ),
                    ),
                ),
            ),
        ),
        profile().model_copy(update={"config": config}),
        {},
    )
    text = "\n".join(p.text for p in Document(BytesIO(rendered)).paragraphs)
    assert "(a)" in text
    assert "(i)" in text
    assert "(ii)" in text
    assert "(b)" in text
    # Verify that (b) appears after (ii), not as a continuation of depth-2
    assert text.index("(b)") > text.index("(ii)")


@pytest.mark.parametrize(
    "style", ["1", "1.", "(1)", "arabic-dot", "lower-alpha", "upper-alpha", "roman"]
)
def test_preview_and_docx_question_labels_match(style: str) -> None:
    numbered = number_questions(
        [
            SectionForNumbering(
                position=1,
                questions=[QuestionForNumbering(question_id="q", position=1, marks=Decimal(1))],
            )
        ],
        NumberingConfig(question_style=style),  # type: ignore[arg-type]
    )
    config = profile().config.model_copy(
        update={"numbering_config_json": TemplateNumberingConfig(question_style=style)}
    )  # type: ignore[arg-type]
    rendered = render_paper(
        paper(),
        profile().model_copy(update={"config": config}),
        {IMAGE_ID: (FIXTURES / "tiny.png").read_bytes()},
    )
    first_question_text = "\n".join(p.text for p in Document(BytesIO(rendered)).paragraphs)
    assert f"{numbered[0].label} 保留原文" in first_question_text


@pytest.mark.parametrize(
    "first_child",
    [
        {
            "type": "bulletList",
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        {
                            "type": "paragraph",
                            "content": [{"type": "text", "text": "List first"}],
                        }
                    ],
                }
            ],
        },
        {
            "type": "table",
            "content": [
                {
                    "type": "tableRow",
                    "content": [
                        {
                            "type": "tableCell",
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": "Table first"}],
                                }
                            ],
                        }
                    ],
                }
            ],
        },
    ],
)
def test_sub_question_can_start_with_list_or_table(first_child: object) -> None:
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "subQuestion",
                    "attrs": {"label": "(a)"},
                    "content": [first_child],
                }
            ],
        }
    )
    extracted = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(_render_content(content))).paragraphs
    )
    assert "(a)" in extracted


def test_subquestion_leaf_marks_render() -> None:
    content = DocNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "subQuestion",
                    "attrs": {"label": "(a)", "marks": "2"},
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "part a"}]}
                    ],
                }
            ],
        }
    )
    extracted = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(_render_content(content))).paragraphs
    )
    assert "(2 marks)" in extracted
