"""Document Parsing layer tests.

Covers the ``DocumentBlock[]`` contract across DOCX and PDF: reading order,
block kinds (text/heading/table/image/equation/caption/header/footer), page
tagging, missing bbox, the scanned/OCR boundary, and malformed input.
"""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml

from app.exam.parsing import (
    BlockKind,
    DocxParser,
    ParseError,
    PdfParser,
    UnsupportedFormatError,
    parse_document,
)

FIXTURES = Path(__file__).parent / "fixtures"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# --- fixtures -------------------------------------------------------------


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Section A", level=1)
    doc.add_paragraph("Question 1 body.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "h1"
    table.cell(0, 1).text = "h2"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    doc.add_paragraph("Figure 1: a diagram")
    doc.add_paragraph("A paragraph with an image:")
    doc.add_picture(str(FIXTURES / "tiny.png"))
    p = doc.add_paragraph()
    p._p.append(parse_xml(f'<m:oMath xmlns:m="{MATH_NS}"><m:r><m:t>x = 2</m:t></m:r></m:oMath>'))
    doc.sections[0].header.paragraphs[0].text = "School Name"
    doc.sections[0].footer.paragraphs[0].text = "Page footer"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(page_texts: list[str | None]) -> bytes:
    """Build a minimal valid PDF with one page per entry (None = blank page)."""
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    n = len(page_texts)
    font_obj = 3 + 2 * n

    def emit(num: int, body: bytes) -> None:
        nonlocal out
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + body + b"endobj\n"

    page_objs = [3 + 2 * i for i in range(n)]
    content_objs = [4 + 2 * i for i in range(n)]
    kids = " ".join(f"{p} 0 R" for p in page_objs)
    emit(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    emit(2, f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    for i, text in enumerate(page_texts):
        page_ref = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_objs[i]} 0 R "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> >>"
        ).encode()
        emit(page_objs[i], page_ref)
        stream = b"" if text is None else f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        emit(
            content_objs[i],
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        )
    emit(font_obj, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    xref_offset = len(out)
    out += f"xref\n0 {font_obj + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for num in range(1, font_obj + 1):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {font_obj + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


# --- DOCX -----------------------------------------------------------------


def test_docx_reading_order_and_kinds() -> None:
    result = DocxParser().parse(_docx_bytes(), source_name="paper.docx")
    kinds = [b.kind for b in result.blocks]
    assert kinds[:3] == [BlockKind.HEADING, BlockKind.TEXT, BlockKind.TABLE]
    assert BlockKind.CAPTION in kinds
    assert BlockKind.IMAGE in kinds
    assert BlockKind.EQUATION in kinds
    # order is strictly increasing and dense
    orders = [b.order for b in result.blocks]
    assert orders == list(range(len(orders)))


def test_docx_heading_level() -> None:
    result = DocxParser().parse(_docx_bytes())
    heading = next(b for b in result.blocks if b.kind == BlockKind.HEADING)
    assert heading.text == "Section A"
    assert heading.meta["level"] == 1


def test_docx_table_rows() -> None:
    result = DocxParser().parse(_docx_bytes())
    table = next(b for b in result.blocks if b.kind == BlockKind.TABLE)
    assert table.rows == [["h1", "h2"], ["a", "b"]]


def test_docx_multiline_cell_preserved() -> None:
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    cell.text = "line1"
    cell.add_paragraph("line2")
    buf = BytesIO()
    doc.save(buf)
    result = DocxParser().parse(buf.getvalue())
    assert result.blocks[0].rows == [["line1\nline2"]]


def test_docx_nested_table_flattened() -> None:
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    nested = parse_xml(
        f'<w:tbl xmlns:w="{W_NS}">'
        "<w:tr><w:tc><w:p><w:r><w:t>構 造</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>風媒花</w:t></w:r></w:p></w:tc></w:tr>"
        "<w:tr><w:tc><w:p><w:r><w:t>花瓣</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>細小</w:t></w:r></w:p></w:tc></w:tr>"
        "</w:tbl>"
    )
    cell._tc.append(nested)
    buf = BytesIO()
    doc.save(buf)
    result = DocxParser().parse(buf.getvalue())
    assert result.blocks[0].rows == [["構 造 | 風媒花\n花瓣 | 細小"]]


def test_docx_drawing_cell_flagged_non_text() -> None:
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    drawing = parse_xml(
        f'<w:p xmlns:w="{W_NS}"><w:r><w:drawing>'
        "<w:txbxContent><w:p><w:r><w:t>garbage textbox text</w:t></w:r></w:p></w:txbxContent>"
        "</w:drawing></w:r></w:p>"
    )
    cell._tc.append(drawing)
    buf = BytesIO()
    doc.save(buf)
    result = DocxParser().parse(buf.getvalue())
    table = result.blocks[0]
    assert table.rows == [[""]]  # no garbled text
    assert table.meta["non_text_cells"] == {"0:0": "drawing"}


def test_docx_image_extracts_asset() -> None:
    result = DocxParser().parse(_docx_bytes())
    image = next(b for b in result.blocks if b.kind == BlockKind.IMAGE)
    assert image.asset is not None
    assert image.asset.mime_type == "image/png"
    assert image.asset.local_id in result.assets
    assert result.assets[image.asset.local_id][:4] == b"\x89PNG"


def test_docx_equation() -> None:
    result = DocxParser().parse(_docx_bytes())
    equation = next(b for b in result.blocks if b.kind == BlockKind.EQUATION)
    assert equation.text == "x = 2"


def test_docx_header_footer() -> None:
    result = DocxParser().parse(_docx_bytes())
    assert any(b.kind == BlockKind.HEADER and b.text == "School Name" for b in result.blocks)
    assert any(b.kind == BlockKind.FOOTER and b.text == "Page footer" for b in result.blocks)


def test_docx_missing_page_and_bbox() -> None:
    result = DocxParser().parse(_docx_bytes())
    assert result.page_count is None
    assert all(b.page is None for b in result.blocks)
    assert all(b.bbox is None for b in result.blocks)


def test_docx_malformed_input() -> None:
    with pytest.raises(ParseError):
        DocxParser().parse(b"this is not a zip file", source_name="x.docx")


# --- PDF ------------------------------------------------------------------


def test_pdf_simple_text_and_pages() -> None:
    result = PdfParser().parse(_pdf_bytes(["Hello page one", "Hello page two"]))
    assert result.page_count == 2
    texts = [b.text for b in result.blocks]
    assert texts == ["Hello page one", "Hello page two"]
    assert [b.page for b in result.blocks] == [1, 2]
    assert all(b.bbox is None for b in result.blocks)  # missing bbox


def test_pdf_scanned_page_flags_review() -> None:
    result = PdfParser().parse(_pdf_bytes([None]))
    assert result.blocks == []
    assert result.needs_review is True
    assert any("no text layer" in w for w in result.warnings)


def test_pdf_malformed_input() -> None:
    with pytest.raises(ParseError):
        PdfParser().parse(b"not a pdf", source_name="x.pdf")


# --- dispatch -------------------------------------------------------------


def test_parse_document_dispatches_by_format() -> None:
    docx = parse_document(_docx_bytes(), filename="p.docx")
    assert docx.format == "docx"
    pdf = parse_document(_pdf_bytes(["hi"]), filename="p.pdf")
    assert pdf.format == "pdf"


def test_parse_document_unsupported() -> None:
    with pytest.raises(UnsupportedFormatError):
        parse_document(b"garbage", filename="x.xyz")
