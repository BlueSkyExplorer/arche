from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument

from app.services.docx_images import extract_cell_images

PNG = (Path(__file__).parent / "fixtures" / "tiny.png").read_bytes()


def test_extract_cell_images_returns_positioned_blobs() -> None:
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 1).add_paragraph().add_run().add_picture(BytesIO(PNG))
    buf = BytesIO()
    doc.save(buf)
    doc2 = DocxDocument(buf)

    found = extract_cell_images(doc2)
    assert len(found) == 1
    (ti, ri, ci), blob = list(found.items())[0]
    assert (ti, ri, ci) == (0, 0, 1)
    assert blob.startswith(b"\x89PNG")
